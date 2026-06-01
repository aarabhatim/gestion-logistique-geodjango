"""
dm_utils/export_excel.py
Fonctions d'export Excel (.xlsx) pour l'admin DeliverMap
"""
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter


# ── Palette couleurs ──────────────────────────────────────────────────────────
DARK_BG    = "1E1E2E"
ACCENT     = "6366F1"   # indigo
HEADER_FG  = "FFFFFF"
ROW_EVEN   = "2A2A3E"
ROW_ODD    = "252538"
TEXT_WHITE = "FFFFFF"
TEXT_MUTED = "A1A1AA"
SUCCESS    = "10B981"
WARNING    = "F59E0B"
DANGER     = "EF4444"


def _thin_border():
    s = Side(style='thin', color="3F3F5A")
    return Border(left=s, right=s, top=s, bottom=s)


def _header_style(ws, row, columns):
    """Style the header row."""
    for col_idx, (col_title, col_width) in enumerate(columns, 1):
        cell = ws.cell(row=row, column=col_idx, value=col_title)
        cell.font = Font(bold=True, color=HEADER_FG, size=11, name="Calibri")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width
    ws.row_dimensions[row].height = 32


def _row_style(ws, row, data, row_num):
    """Style a data row."""
    bg = ROW_EVEN if row_num % 2 == 0 else ROW_ODD
    for col_idx, value in enumerate(data, 1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(color=TEXT_WHITE, size=10, name="Calibri")
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        cell.border = _thin_border()
    ws.row_dimensions[row].height = 22


def _title_row(ws, title, ncols):
    """Merge first row as title banner."""
    ws.row_dimensions[1].height = 40
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title_cell = ws.cell(row=1, column=1, value=f"📦 DeliverMap — {title}")
    title_cell.font = Font(bold=True, color=HEADER_FG, size=14, name="Calibri")
    title_cell.fill = PatternFill("solid", fgColor=DARK_BG)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sub row: export date
    ws.row_dimensions[2].height = 20
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    date_cell = ws.cell(row=2, column=1, value=f"Exporté le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    date_cell.font = Font(color=TEXT_MUTED, size=9, name="Calibri", italic=True)
    date_cell.fill = PatternFill("solid", fgColor=DARK_BG)
    date_cell.alignment = Alignment(horizontal="center", vertical="center")


def _finalize(wb, ws):
    """Freeze header, set tab colour, return bytes."""
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = ACCENT
    ws.sheet_view.showGridLines = False

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── COMMANDES ─────────────────────────────────────────────────────────────────
def export_commandes_xlsx(queryset):
    """Exporte le queryset Commande en .xlsx, retourne bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commandes"
    ws.sheet_view.showGridLines = False

    columns = [
        ("Référence",          14),
        ("Date",               18),
        ("Client",             20),
        ("Boutique",           22),
        ("Transporteur",       20),
        ("Statut",             16),
        ("Paiement",           14),
        ("Payée",              10),
        ("Sous-total (MAD)",   16),
        ("Frais livraison",    15),
        ("Réduction",          12),
        ("Total (MAD)",        14),
        ("Adresse livraison",  35),
    ]

    _title_row(ws, "Export Commandes", len(columns))
    _header_style(ws, 3, columns)

    STATUT_LABELS = {
        'EN_ATTENTE':    'En attente',
        'VALIDEE':       'Validée',
        'EN_PREPARATION':'En préparation',
        'EN_ROUTE':      'En route',
        'LIVREE':        'Livrée',
        'ANNULEE':       'Annulée',
    }

    for i, cmd in enumerate(queryset, 1):
        transporteur = cmd.transporteur.get_full_name() if cmd.transporteur else '—'
        row_data = [
            cmd.reference,
            cmd.created_at.strftime('%d/%m/%Y %H:%M') if cmd.created_at else '',
            cmd.client.get_full_name() or cmd.client.username,
            cmd.fondateur.nom_boutique if cmd.fondateur else '—',
            transporteur,
            STATUT_LABELS.get(cmd.statut, cmd.statut),
            cmd.get_mode_paiement_display(),
            'Oui' if cmd.est_paye else 'Non',
            float(cmd.sous_total),
            float(cmd.frais_livraison),
            float(cmd.reduction),
            float(cmd.total_price),
            cmd.adresse_livraison,
        ]
        _row_style(ws, i + 3, row_data, i)

        # Couleur statut colonne F (col 6)
        status_cell = ws.cell(row=i + 3, column=6)
        color_map = {
            'LIVREE': SUCCESS, 'ANNULEE': DANGER,
            'EN_ROUTE': WARNING, 'EN_PREPARATION': "8B5CF6",
        }
        c = color_map.get(cmd.statut)
        if c:
            status_cell.font = Font(color=c, bold=True, size=10, name="Calibri")

    return _finalize(wb, ws)


# ── TRANSPORTEURS ─────────────────────────────────────────────────────────────
def export_transporteurs_xlsx(queryset):
    """Exporte le queryset Transporteur en .xlsx, retourne bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transporteurs"

    columns = [
        ("Nom complet",        22),
        ("Nom d'utilisateur",  18),
        ("Email",              26),
        ("Téléphone",          16),
        ("Véhicule",           14),
        ("Plaque",             14),
        ("Capacité (kg)",      14),
        ("Vérifié",            10),
        ("Disponible",         12),
        ("Note moyenne",       13),
        ("Nb avis",            10),
        ("Nb livraisons",      14),
        ("Revenus total (MAD)",18),
        ("Date inscription",   18),
    ]

    _title_row(ws, "Export Transporteurs", len(columns))
    _header_style(ws, 3, columns)

    for i, t in enumerate(queryset, 1):
        row_data = [
            t.user.get_full_name(),
            t.user.username,
            t.user.email,
            getattr(t.user, 'telephone', '') or '',
            t.get_vehicule_type_display(),
            t.plaque,
            t.capacite_kg,
            'Oui' if t.is_verified else 'Non',
            'Oui' if t.is_available else 'Non',
            t.note_moyenne,
            t.nombre_avis,
            t.nombre_livraisons,
            float(t.revenus_total),
            t.date_inscription.strftime('%d/%m/%Y') if t.date_inscription else '',
        ]
        _row_style(ws, i + 3, row_data, i)

        # Couleur vérifié colonne H (col 8)
        ver_cell = ws.cell(row=i + 3, column=8)
        ver_cell.font = Font(
            color=SUCCESS if t.is_verified else DANGER,
            bold=True, size=10, name="Calibri"
        )

    return _finalize(wb, ws)


# ── CLIENTS ───────────────────────────────────────────────────────────────────
def export_clients_xlsx(queryset):
    """Exporte le queryset Client en .xlsx, retourne bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients"

    columns = [
        ("Prénom",             16),
        ("Nom",                16),
        ("Email",              26),
        ("Téléphone",          16),
        ("Entreprise",         20),
        ("Adresse",            30),
        ("Note fidélité",      14),
        ("Actif",              10),
        ("Date inscription",   18),
    ]

    _title_row(ws, "Export Clients", len(columns))
    _header_style(ws, 3, columns)

    for i, c in enumerate(queryset, 1):
        row_data = [
            c.prenom,
            c.nom,
            c.email,
            c.telephone,
            c.entreprise or '—',
            c.adresse,
            c.note_fidelite,
            'Oui' if c.actif else 'Non',
            c.date_inscription.strftime('%d/%m/%Y') if c.date_inscription else '',
        ]
        _row_style(ws, i + 3, row_data, i)

        # Couleur actif colonne H (col 8)
        actif_cell = ws.cell(row=i + 3, column=8)
        actif_cell.font = Font(
            color=SUCCESS if c.actif else DANGER,
            bold=True, size=10, name="Calibri"
        )

    return _finalize(wb, ws)
